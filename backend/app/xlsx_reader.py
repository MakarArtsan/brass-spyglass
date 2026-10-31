"""Чтение Excel-трекера в нейтральную структуру. Ничего не пишет в БД.

Колонки ищутся по заголовкам, не по буквам (ТЗ, «Импорт Excel-трекера»). Формульные колонки
(таймкоды, суммы, % готовности) не читаются: на сайте они считаются заново в metrics.py.
Всё подозрительное попадает в `Bundle.warnings` — импорт из-за этого не падает.
"""

import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from app.constants import (
    ASSET_TYPE_EXCEL_LABELS,
    PRIORITY_EXCEL_LABELS,
    SHOT_SIZE_EXCEL_LABELS,
    SOUND_TYPE_EXCEL_LABELS,
    STAGE_EXCEL_LABELS,
    STATUS_EXCEL_LABELS,
    TASK_BLOCK_EXCEL_LABELS,
    TOOL_EXCEL_LABELS,
    AssetType,
    Priority,
    ShotSize,
    SoundType,
    Stage,
    Status,
    TaskBlock,
    Tool,
    from_excel,
)

NONE_MARKS = {"—", "–", "-", ""}


@dataclass
class SceneRow:
    number: int
    title: str
    location: str | None
    time_of_day: str | None
    planned_sec: float
    characters: str | None
    key_assets: str | None
    lines_en: str | None
    drum_level: int | None
    meaning: str | None
    cut_notes: str | None
    status: str
    priority: str | None
    due_date: date | None
    comment: str | None


@dataclass
class ShotRow:
    code: str
    scene_number: int
    description: str | None
    size: str | None
    camera: str | None
    duration_sec: float
    asset_codes: list[str]
    sound_note: str | None
    tool: str | None
    stage: str | None
    takes_count: int
    status: str
    priority: str | None
    due_date: date | None
    comment: str | None


@dataclass
class AssetRow:
    code: str
    type: str
    name: str
    description: str | None
    states: list[str]
    status: str
    priority: str | None
    due_date: date | None
    comment: str | None


@dataclass
class SetupRow:
    number: int
    motif: str
    setup_codes: list[str]
    payoff_codes: list[str]
    reader_takeaway: str | None
    risk: str | None
    status: str
    comment: str | None


@dataclass
class SoundRow:
    code: str
    scene_number: int | None
    shot_code: str | None
    type: str
    who: str | None
    text_en: str | None
    text_ru: str | None
    rights: str | None
    status: str
    comment: str | None


@dataclass
class TaskRow:
    code: str
    block: str | None
    title: str
    due_date: date | None
    status: str
    priority: str | None
    comment: str | None


@dataclass
class Bundle:
    scenes: list[SceneRow] = field(default_factory=list)
    shots: list[ShotRow] = field(default_factory=list)
    assets: list[AssetRow] = field(default_factory=list)
    setups: list[SetupRow] = field(default_factory=list)
    sounds: list[SoundRow] = field(default_factory=list)
    tasks: list[TaskRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class WorkbookError(Exception):
    """В файле нет листа или обязательной колонки — продолжать нельзя."""


# --- значения ячеек ---------------------------------------------------------------------------


def text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    cleaned = re.sub(r"[ \t]+", " ", str(value)).strip()
    return None if cleaned in NONE_MARKS else cleaned


def number(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(str(value).replace(",", ".").strip())
    except ValueError:
        return None


def integer(value: Any) -> int | None:
    n = number(value)
    return None if n is None else int(n)


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return datetime.fromisoformat(value.strip()).date()
        except ValueError:
            return None
    return None


def split_states(value: Any) -> list[str]:
    """«1) белая 2) в саже» и «Дозор / таверна / финал» → отдельные состояния ассета."""
    raw = text(value)
    if raw is None:
        return []
    if re.match(r"^\s*1\)", raw):
        parts = re.split(r"(?:^|\s)\d+\)\s*", raw)
    else:
        parts = re.split(r"\s*/\s*", raw)
    return [p.strip(" ;") for p in parts if p.strip(" ;")]


def split_asset_codes(value: Any) -> list[str]:
    raw = text(value)
    if raw is None:
        return []
    seen: list[str] = []
    for token in re.split(r"[,;\s]+", raw):
        if token and token not in seen:
            seen.append(token)
    return seen


_RANGE = re.compile(r"^(\d{1,2}\.\d{1,3})\s*[–—-]\s*(\d{1,2}\.\d{1,3})$")


def _norm_code(code: str) -> str:
    scene, shot = code.split(".")
    return f"{int(scene):02d}.{int(shot):02d}"


def expand_shot_refs(
    value: Any, ordered_codes: list[str], warnings: list[str], where: str
) -> list[str]:
    """«01.04, 11.03–11.05, 10.01–11.02» → список кодов шотов. Диапазон идёт по порядку шотов
    в листе «Шоты», поэтому работает и через границу сцен."""
    raw = text(value)
    if raw is None:
        return []
    position = {c: i for i, c in enumerate(ordered_codes)}
    out: list[str] = []

    def add(code: str) -> None:
        if code not in out:
            out.append(code)

    for token in re.split(r"\s*[,;]\s*", raw):
        token = token.strip()
        if not token:
            continue
        if match := _RANGE.match(token):
            start, end = _norm_code(match.group(1)), _norm_code(match.group(2))
            if start in position and end in position and position[start] <= position[end]:
                for code in ordered_codes[position[start] : position[end] + 1]:
                    add(code)
            else:
                warnings.append(f"{where}: диапазон «{token}» не найден в листе «Шоты»")
        elif re.fullmatch(r"\d{1,2}\.\d{1,3}", token):
            code = _norm_code(token)
            if code in position:
                add(code)
            else:
                warnings.append(f"{where}: шот «{token}» не найден в листе «Шоты»")
        else:
            warnings.append(f"{where}: не удалось разобрать «{token}»")
    return out


# --- листы -----------------------------------------------------------------------------------


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _rows(wb: openpyxl.Workbook, sheet: str, columns: dict[str, tuple[str, bool]]) -> list[dict]:
    """Строки листа как словари. `columns`: поле → (заголовок или его начало, обязательна ли)."""
    if sheet not in wb.sheetnames:
        raise WorkbookError(f"В файле нет листа «{sheet}»")
    ws = wb[sheet]
    headers = [_norm_header(c.value) for c in ws[1]]
    index: dict[str, int | None] = {}
    for key, (title, required) in columns.items():
        found = next((i for i, h in enumerate(headers) if h == title), None)
        if found is None:
            found = next((i for i, h in enumerate(headers) if h.startswith(title)), None)
        if found is None and required:
            raise WorkbookError(f"На листе «{sheet}» нет колонки «{title}»")
        index[key] = found
    first_key = next(iter(columns))
    rows = []
    for row_no, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        pos = index[first_key]
        if pos is None or pos >= len(cells) or text(cells[pos]) is None:
            continue
        item = {
            k: (cells[i] if i is not None and i < len(cells) else None) for k, i in index.items()
        }
        item["_row"] = row_no
        rows.append(item)
    return rows


def _lookup(enum_cls: Any, labels: dict, value: Any, where: str, warnings: list[str]) -> str | None:
    label = text(value)
    if label is None:
        return None
    member = from_excel(enum_cls, labels, label)
    if member is None:
        warnings.append(f"{where}: неизвестное значение «{label}»")
        return None
    return member.value


def _status(value: Any, where: str, warnings: list[str]) -> str:
    found = _lookup(Status, STATUS_EXCEL_LABELS, value, where, warnings)
    return found or Status.NOT_STARTED.value


def _unmapped_assignee(row: dict, where: str, warnings: list[str]) -> None:
    if text(row.get("assignee")) is not None:
        warnings.append(
            f"{where}: ответственный «{text(row['assignee'])}» не сопоставлен с участником"
        )


def read_workbook(path: str | Path) -> Bundle:
    with warnings.catch_warnings():
        # Условное форматирование и расширения Excel нам не нужны: openpyxl лишь предупреждает.
        warnings.filterwarnings(
            "ignore", message=".*extension is not supported.*", category=UserWarning
        )
        wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    bundle = Bundle()
    w = bundle.warnings
    common = {
        "status": ("Статус", False),
        "priority": ("Приоритет", False),
        "assignee": ("Ответственный", False),
        "due": ("Дедлайн", False),
        "comment": ("Комментарий", False),
    }

    for r in _rows(
        wb,
        "Сцены",
        {
            "number": ("№", True),
            "title": ("Сцена", True),
            "location": ("Локация", False),
            "time_of_day": ("Время суток", False),
            "planned": ("План, сек", True),
            "characters": ("Персонажи", False),
            "key_assets": ("Ключевые ассеты", False),
            "lines_en": ("Реплики (EN)", False),
            "drum": ("Барабан", False),
            "meaning": ("Смысл сцены", False),
            "cut_notes": ("Если сокращать", False),
            **common,
        },
    ):
        where = f"Сцены, строка {r['_row']}"
        _unmapped_assignee(r, where, w)
        number_ = integer(r["number"])
        if number_ is None:
            w.append(f"{where}: нет номера сцены, строка пропущена")
            continue
        bundle.scenes.append(
            SceneRow(
                number=number_,
                title=text(r["title"]) or f"Сцена {number_}",
                location=text(r["location"]),
                time_of_day=text(r["time_of_day"]),
                planned_sec=number(r["planned"]) or 0.0,
                characters=text(r["characters"]),
                key_assets=text(r["key_assets"]),
                lines_en=text(r["lines_en"]),
                drum_level=integer(r["drum"]),
                meaning=text(r["meaning"]),
                cut_notes=text(r["cut_notes"]),
                status=_status(r["status"], where, w),
                priority=_lookup(Priority, PRIORITY_EXCEL_LABELS, r["priority"], where, w),
                due_date=as_date(r["due"]),
                comment=text(r["comment"]),
            )
        )

    shot_rows = _rows(
        wb,
        "Шоты",
        {
            "code": ("ID", True),
            "scene": ("Сцена", True),
            "description": ("Описание кадра", False),
            "size": ("Крупность", False),
            "camera": ("Камера", False),
            "duration": ("Длит.", True),
            "assets": ("Ассеты в кадре", False),
            "sound": ("Реплика / звук", False),
            "tool": ("Инструмент", False),
            "stage": ("Этап", False),
            "takes": ("Дублей", False),
            "selected": ("Выбранный дубль", False),
            **common,
        },
    )
    for r in shot_rows:
        where = f"Шоты, строка {r['_row']}"
        _unmapped_assignee(r, where, w)
        code = text(r["code"])
        scene_no = integer(r["scene"])
        if code is None or scene_no is None or not re.fullmatch(r"\d{1,2}\.\d{1,3}", code):
            w.append(f"{where}: не разобран ID шота или номер сцены, строка пропущена")
            continue
        if text(r["selected"]) is not None:
            w.append(
                f"{where}: «Выбранный дубль / файл» не импортируется — файл прикрепляют на сайте"
            )
        bundle.shots.append(
            ShotRow(
                code=_norm_code(code),
                scene_number=scene_no,
                description=text(r["description"]),
                size=_lookup(ShotSize, SHOT_SIZE_EXCEL_LABELS, r["size"], where, w),
                camera=text(r["camera"]),
                duration_sec=number(r["duration"]) or 0.0,
                asset_codes=split_asset_codes(r["assets"]),
                sound_note=text(r["sound"]),
                tool=_lookup(Tool, TOOL_EXCEL_LABELS, r["tool"], where, w),
                stage=_lookup(Stage, STAGE_EXCEL_LABELS, r["stage"], where, w),
                takes_count=integer(r["takes"]) or 0,
                status=_status(r["status"], where, w),
                priority=_lookup(Priority, PRIORITY_EXCEL_LABELS, r["priority"], where, w),
                due_date=as_date(r["due"]),
                comment=text(r["comment"]),
            )
        )
    ordered_codes = [s.code for s in bundle.shots]

    for r in _rows(
        wb,
        "Ассеты",
        {
            "code": ("ID", True),
            "type": ("Тип", True),
            "name": ("Название", True),
            "description": ("Описание", False),
            "states": ("Состояния и версии", False),
            "reference": ("Референс", False),
            **common,
        },
    ):
        where = f"Ассеты, строка {r['_row']}"
        _unmapped_assignee(r, where, w)
        if text(r["reference"]) is not None:
            w.append(
                f"{where}: «Референс (ссылка / файл)» не импортируется — добавьте файл на сайте"
            )
        asset_type = _lookup(AssetType, ASSET_TYPE_EXCEL_LABELS, r["type"], where, w)
        bundle.assets.append(
            AssetRow(
                code=text(r["code"]) or "",
                type=asset_type or AssetType.PROP.value,
                name=text(r["name"]) or text(r["code"]) or "",
                description=text(r["description"]),
                states=split_states(r["states"]),
                status=_status(r["status"], where, w),
                priority=_lookup(Priority, PRIORITY_EXCEL_LABELS, r["priority"], where, w),
                due_date=as_date(r["due"]),
                comment=text(r["comment"]),
            )
        )

    known_assets = {a.code for a in bundle.assets}
    for shot in bundle.shots:
        for code in shot.asset_codes:
            if code not in known_assets:
                w.append(f"Шот {shot.code}: ассет «{code}» не найден в листе «Ассеты»")

    for r in _rows(
        wb,
        "Сетапы",
        {
            "number": ("№", True),
            "motif": ("Мотив", True),
            "setup": ("Сетап (шот)", False),
            "payoff": ("Пэйофф (шот)", False),
            "takeaway": ("Что должен считать зритель", False),
            "risk": ("Риск", False),
            "status": ("Статус", False),
            "comment": ("Комментарий", False),
        },
    ):
        where = f"Сетапы, строка {r['_row']}"
        number_ = integer(r["number"])
        if number_ is None:
            w.append(f"{where}: нет номера, строка пропущена")
            continue
        bundle.setups.append(
            SetupRow(
                number=number_,
                motif=text(r["motif"]) or f"Сетап {number_}",
                setup_codes=expand_shot_refs(r["setup"], ordered_codes, w, f"{where}, сетап"),
                payoff_codes=expand_shot_refs(r["payoff"], ordered_codes, w, f"{where}, пэйофф"),
                reader_takeaway=text(r["takeaway"]),
                risk=text(r["risk"]),
                status=_status(r["status"], where, w),
                comment=text(r["comment"]),
            )
        )

    for r in _rows(
        wb,
        "Звук",
        {
            "code": ("ID", True),
            "scene": ("Сцена", False),
            "shot": ("Шот", False),
            "type": ("Тип", True),
            "who": ("Кто", False),
            "text_en": ("Текст EN", False),
            "text_ru": ("Перевод RU", False),
            "rights": ("Озвучка", False),
            "status": ("Статус", False),
            "comment": ("Комментарий", False),
        },
    ):
        where = f"Звук, строка {r['_row']}"
        shot_code = text(r["shot"])
        if shot_code and re.fullmatch(r"\d{1,2}\.\d{1,3}", shot_code):
            shot_code = _norm_code(shot_code)
            if shot_code not in ordered_codes:
                w.append(f"{where}: шот «{shot_code}» не найден в листе «Шоты»")
                shot_code = None
        elif shot_code:
            w.append(f"{where}: не удалось разобрать шот «{shot_code}»")
            shot_code = None
        sound_type = _lookup(SoundType, SOUND_TYPE_EXCEL_LABELS, r["type"], where, w)
        bundle.sounds.append(
            SoundRow(
                code=text(r["code"]) or "",
                scene_number=integer(r["scene"]),
                shot_code=shot_code,
                type=sound_type or SoundType.SFX.value,
                who=text(r["who"]),
                text_en=text(r["text_en"]),
                text_ru=text(r["text_ru"]),
                rights=text(r["rights"]),
                status=_status(r["status"], where, w),
                comment=text(r["comment"]),
            )
        )

    for r in _rows(
        wb,
        "Задачи",
        {
            "code": ("ID", True),
            "block": ("Блок", False),
            "title": ("Задача", True),
            **common,
        },
    ):
        where = f"Задачи, строка {r['_row']}"
        _unmapped_assignee(r, where, w)
        bundle.tasks.append(
            TaskRow(
                code=text(r["code"]) or "",
                block=_lookup(TaskBlock, TASK_BLOCK_EXCEL_LABELS, r["block"], where, w),
                title=text(r["title"]) or text(r["code"]) or "",
                due_date=as_date(r["due"]),
                status=_status(r["status"], where, w),
                priority=_lookup(Priority, PRIORITY_EXCEL_LABELS, r["priority"], where, w),
                comment=text(r["comment"]),
            )
        )
    return bundle
